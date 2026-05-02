#!/usr/bin/env python3
"""
Comparison table: RDMA FUSEE vs Old CXL (Propose-Vote-Commit) vs New CXL (LFM).
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

os.makedirs("docs", exist_ok=True)

# ── Colors ──
P_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
P_DG    = RGBColor(0x33, 0x33, 0x33)
P_GRAY  = RGBColor(0x9E, 0x9E, 0x9E)
P_LGRAY = RGBColor(0xF5, 0xF5, 0xF5)

P_RDMA   = RGBColor(0xC6, 0x28, 0x28); P_RDMA_L = RGBColor(0xFF, 0xCD, 0xD2)
P_OLD    = RGBColor(0x6A, 0x1B, 0x9A); P_OLD_L  = RGBColor(0xE1, 0xBE, 0xE7)
P_NEW    = RGBColor(0x2E, 0x7D, 0x32); P_NEW_L  = RGBColor(0xC8, 0xE6, 0xC9)

P_HDR    = RGBColor(0xE6, 0x51, 0x00); P_HDR_L  = RGBColor(0xFF, 0xE0, 0xB2)
P_HI     = RGBColor(0xFF, 0x8F, 0x00)


def _in(v):
    return Inches(v)

def _box(slide, l, t, w, h, fc, ec=None, bw=Pt(1.5), text="", fs=10, c=P_DG, bold=True):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                               _in(l), _in(t), _in(w), _in(h))
    s.fill.solid(); s.fill.fore_color.rgb = fc
    if ec:
        s.line.color.rgb = ec; s.line.width = bw
    else:
        s.line.fill.background()
    if text:
        tf = s.text_frame; tf.word_wrap = True
        tf.margin_top = Pt(4); tf.margin_bottom = Pt(4)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(0); p.space_after = Pt(0)
        run = p.add_run(); run.text = text
        run.font.size = Pt(fs); run.font.color.rgb = c; run.font.bold = bold

def _text(slide, l, t, w, h, text, fs=10, c=P_DG, bold=False, align=PP_ALIGN.CENTER):
    tb = slide.shapes.add_textbox(_in(l), _in(t), _in(w), _in(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    run.font.size = Pt(fs); run.font.color.rgb = c; run.font.bold = bold


prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])

# ── Title ──
_text(slide, 0, 0.25, 13.33, 0.55,
      "Consensus Design Comparison: RDMA vs Old CXL vs New CXL",
      20, P_DG, bold=True)

# ── Layout constants ──
left_margin = 0.5
row_height = 0.85
header_height = 0.65
col_label_w = 2.6
col_data_w = 3.3
col_gap = 0.1

# Column X positions
col_xs = [left_margin,
          left_margin + col_label_w + col_gap,
          left_margin + col_label_w + col_gap + col_data_w + col_gap,
          left_margin + col_label_w + col_gap + 2*(col_data_w + col_gap)]
col_widths = [col_label_w, col_data_w, col_data_w, col_data_w]

# ── Header row ──
header_y = 1.15
headers = [
    ("",                  P_WHITE, P_WHITE),
    ("RDMA FUSEE",        P_RDMA_L, P_RDMA),
    ("Old CXL\n(Propose-Vote-Commit)", P_OLD_L, P_OLD),
    ("New CXL\n(Per-Slot LFM)", P_NEW_L, P_NEW),
]
for x, w, (label, fc, ec) in zip(col_xs, col_widths, headers):
    if label:
        _box(slide, x, header_y, w, header_height, fc, ec, Pt(2.5),
             label, 11, ec, bold=True)

# ── Data rows ──
rows = [
    ("Commit point",
     "Step 6: CAS Primary",
     "Step 5: write status=COMMITTED",
     "Step 5: write SlotLock.value"),
    ("Concurrency control",
     "Hardware CAS atomicity",
     "Majority voting (2 of 3)",
     "LFM software mutex"),
    ("Conflict handling",
     "Check CAS return -> retry",
     "Vote failed -> ABORT + retry",
     "Spin-wait -> auto queue"),
    ("Sync ops (uncontended)",
     "Multiple RDMA round-trips",
     "~10 CXL ops + voter polling wait",
     "~12 CXL ops, no waiting"),
    ("Latency estimate",
     "1-3 us (RDMA hardware)",
     "15-50 us (voter polling)",
     "5-10 us (LFM fast path)"),
    ("Hardware requirement",
     "RDMA NIC + atomic ops",
     "CXL load/store + sfence",
     "CXL load/store + sfence"),
    ("Race in claim phase",
     "N/A (hardware atomic)",
     "Yes (two nodes can both win)",
     "Impossible (LFM is mutex)"),
]

start_y = header_y + header_height + 0.1
for ri, (label, c1, c2, c3) in enumerate(rows):
    y = start_y + ri * row_height
    # Label cell
    _box(slide, col_xs[0], y, col_widths[0], row_height-0.05,
         P_HDR_L, P_HDR, Pt(1), label, 10, P_DG, bold=True)
    # Data cells
    _box(slide, col_xs[1], y, col_widths[1], row_height-0.05,
         P_LGRAY, P_RDMA, Pt(0.8), c1, 9, P_DG, bold=False)
    _box(slide, col_xs[2], y, col_widths[2], row_height-0.05,
         P_LGRAY, P_OLD, Pt(0.8), c2, 9, P_DG, bold=False)
    _box(slide, col_xs[3], y, col_widths[3], row_height-0.05,
         P_NEW_L, P_NEW, Pt(1.2), c3, 9, P_DG, bold=False)

# ── Bottom insight box ──
insight_y = start_y + len(rows) * row_height + 0.1
_box(slide, 0.5, insight_y, 12.3, 0.7,
     RGBColor(0xFF, 0xF8, 0xE1), P_HI, Pt(2))
_text(slide, 0.7, insight_y + 0.05, 11.9, 0.3,
      "Key insight:",
      11, P_HI, bold=True, align=PP_ALIGN.LEFT)
_text(slide, 0.7, insight_y + 0.32, 11.9, 0.4,
      "With a real mutex, no voting is needed - the lock IS the consensus.",
      11, P_DG, bold=True, align=PP_ALIGN.LEFT)


out = "docs/consensus_comparison.pptx"
prs.save(out)
print(f"Saved {out}")
print(f"Slides: {len(prs.slides)}")
