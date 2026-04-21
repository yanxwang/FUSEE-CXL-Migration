#!/usr/bin/env python3
"""
CXL consensus step-by-step PPTX.
Unified nodes (top) -> CXL shared memory (bottom).
One slide per step.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

# ── Colors ──
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
COMP    = RGBColor(0x15, 0x65, 0xC0)  # proposer
COMP_L  = RGBColor(0xBB, 0xDE, 0xFB)
VOTER   = RGBColor(0x5E, 0x35, 0xB1)  # voter
VOTER_L = RGBColor(0xD1, 0xC4, 0xE9)
CXL     = RGBColor(0xE6, 0x51, 0x00)  # cxl memory
CXL_L   = RGBColor(0xFF, 0xE0, 0xB2)
IDX     = RGBColor(0xFF, 0xF9, 0xC4)
DATA    = RGBColor(0xE3, 0xF2, 0xFD)
LOG     = RGBColor(0xF3, 0xE5, 0xF5)
CONS    = RGBColor(0xFF, 0xF3, 0xE0)  # consensus log
GRAY    = RGBColor(0x9E, 0x9E, 0x9E)
DG      = RGBColor(0x33, 0x33, 0x33)
OK      = RGBColor(0x1B, 0x5E, 0x20)
FAIL    = RGBColor(0xB7, 0x1C, 0x1C)
COMMIT  = RGBColor(0xFF, 0x6F, 0x00)
HI      = RGBColor(0xFF, 0x8F, 0x00)
LGRAY   = RGBColor(0xEE, 0xEE, 0xEE)

def _in(v):
    return Inches(v)

def _add_box(slide, left, top, w, h, fill, border_color=None, border_w=Pt(1.5),
             text="", font_size=9, font_color=DG, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   _in(left), _in(top), _in(w), _in(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_w
    else:
        shape.line.fill.background()
    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(font_size)
        run.font.color.rgb = font_color
        run.font.bold = bold
    return shape

def _add_text(slide, left, top, w, h, text, font_size=9, color=DG, bold=False,
              align=PP_ALIGN.CENTER):
    txBox = slide.shapes.add_textbox(_in(left), _in(top), _in(w), _in(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = color
    run.font.bold = bold
    return txBox

def _add_arrow(slide, x0, y0, x1, y1, color=COMP, width=Pt(2), dashed=False):
    connector = slide.shapes.add_connector(1, _in(x0), _in(y0), _in(x1), _in(y1))
    connector.line.color.rgb = color
    connector.line.width = width
    if dashed:
        connector.line.dash_style = 2
    fmt = connector.line._ln
    tail = fmt.makeelement(qn('a:tailEnd'), {})
    tail.set('type', 'triangle')
    tail.set('w', 'med')
    tail.set('len', 'med')
    fmt.append(tail)
    return connector

# ── Layout constants ──
COL_X = [0.5, 4.8, 9.1]  # 3 node columns
NW = 3.8      # node width
NH = 2.0      # node height
NY = 0.8      # node top Y

CXL_X = 0.3   # CXL memory left
CXL_Y = 4.2   # CXL memory top
CXL_W = 12.7  # CXL memory width
CXL_H = 2.8   # CXL memory height

# region sizes inside nodes
RW_IDX = 1.7
RW_KV = 1.7
RH = 0.55
R_Y = 0.55     # relative to node top

# CXL internal region positions
CL_X = CXL_X + 0.2         # consensus log
CL_Y = CXL_Y + 0.5
CL_W = 4.8
CL_H = 1.9

OL_X = CXL_X + 5.3         # oplog
OL_Y = CL_Y
OL_W = 2.5
OL_H = 1.9

ST_X = CXL_X + 8.1         # staging
ST_Y = CL_Y
ST_W = 2.2
ST_H = 1.9

AD_X = CXL_X + 10.6        # alloc dir
AD_Y = CL_Y
AD_W = 1.8
AD_H = 1.9


def _draw_node(slide, col, title, nc, ncl, idx_text="Slot X = 0x0", kv_text="(empty)",
               hi_idx=False, hi_kv=False):
    x = COL_X[col]
    _add_box(slide, x, NY, NW, NH, ncl, nc, Pt(2), title, 10, nc)
    # Index
    bc = HI if hi_idx else GRAY
    bw = Pt(2.5) if hi_idx else Pt(0.8)
    fc = OK if hi_idx else DG
    _add_box(slide, x+0.1, NY+R_Y, RW_IDX, RH, IDX, bc, bw,
             f"Local Index\n{idx_text}", 7, fc, bold=hi_idx)
    # KV Data
    bc = HI if hi_kv else GRAY
    bw = Pt(2.5) if hi_kv else Pt(0.8)
    fc = OK if hi_kv else DG
    _add_box(slide, x+1.95, NY+R_Y, RW_KV, RH, DATA, bc, bw,
             f"Local KV\n{kv_text}", 7, fc, bold=hi_kv)
    # Voter fiber bar
    _add_box(slide, x+0.1, NY+1.3, NW-0.2, 0.4, LGRAY, GRAY, Pt(0.5),
             "Voter Fiber (polls CXL)", 6.5, GRAY)


def _draw_cxl(slide, cl_lines=None, ol_text="op_status = ?", st_text="(empty)",
              hi_cl=False, hi_ol=False, hi_st=False):
    # main box
    _add_box(slide, CXL_X, CXL_Y, CXL_W, CXL_H, CXL_L, CXL, Pt(2.5),
             "", 10, CXL)
    _add_text(slide, CXL_X, CXL_Y+0.05, CXL_W, 0.35,
              "CXL Type 3 Shared Memory", 12, CXL, bold=True)

    # ConsensusLog
    bc = HI if hi_cl else GRAY
    bw = Pt(2.5) if hi_cl else Pt(0.8)
    _add_box(slide, CL_X, CL_Y, CL_W, CL_H, CONS, bc, bw, "", 8, DG)
    _add_text(slide, CL_X, CL_Y+0.02, CL_W, 0.3,
              "ConsensusLog Entry", 9, CXL, bold=True)
    if cl_lines:
        for i, line in enumerate(cl_lines):
            _add_text(slide, CL_X+0.1, CL_Y+0.3+i*0.3, CL_W-0.2, 0.28,
                      line, 7.5, DG, bold=False, align=PP_ALIGN.LEFT)

    # OpLog
    bc = HI if hi_ol else GRAY
    bw = Pt(2.5) if hi_ol else Pt(0.8)
    _add_box(slide, OL_X, OL_Y, OL_W, OL_H, LOG, bc, bw,
             f"OpLog\n{ol_text}", 8, OK if hi_ol else DG, bold=hi_ol)

    # Staging
    bc = HI if hi_st else GRAY
    bw = Pt(2.5) if hi_st else Pt(0.8)
    _add_box(slide, ST_X, ST_Y, ST_W, ST_H, DATA, bc, bw,
             f"Staging\n{st_text}", 8, OK if hi_st else DG, bold=hi_st)

    # Alloc Dir
    _add_box(slide, AD_X, AD_Y, AD_W, AD_H, LGRAY, GRAY, Pt(0.8),
             "Alloc\nDirectory", 7, GRAY)


def _add_slide(title_text, title_color=COMP):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _add_text(slide, 0, 0.05, 13.33, 0.5, title_text, 16, title_color, bold=True)
    return slide


def _legend(slide):
    y = 7.1
    items = [("Proposer", COMP_L, COMP), ("Voter", VOTER_L, VOTER),
             ("CXL Mem", CXL_L, CXL), ("Index", IDX, GRAY),
             ("ConsensusLog", CONS, GRAY), ("OpLog", LOG, GRAY),
             ("KV/Staging", DATA, GRAY), ("Modified", WHITE, HI)]
    for i, (lbl, fc, ec) in enumerate(items):
        x = 0.8 + i * 1.55
        _add_box(slide, x, y, 0.25, 0.18, fc, ec, Pt(1.5))
        _add_text(slide, x+0.3, y-0.02, 1.1, 0.22, lbl, 7, DG)


def _switch_lines(slide):
    """Dashed lines from nodes to CXL."""
    for col in range(3):
        x = COL_X[col] + NW/2
        _add_arrow(slide, x, NY+NH, x, CXL_Y, GRAY, Pt(0.8), dashed=True)
    _add_text(slide, 5.5, 3.45, 2, 0.25, "CXL Switch", 8, GRAY)


# default consensus log (empty)
CL_EMPTY = [
    "proposer = ?    epoch = ?",
    "old_slot = ?    new_slot = ?",
    "vote[0]=?  vote[1]=?  vote[2]=?",
    "status = FREE",
]

CL_PROPOSED = [
    "proposer = 0    epoch = 42",
    "old_slot = 0x0    new_slot = A",
    "vote[0]=ACK  vote[1]=?  vote[2]=?",
    "status = PROPOSED",
]

CL_VOTED = [
    "proposer = 0    epoch = 42",
    "old_slot = 0x0    new_slot = A",
    "vote[0]=ACK  vote[1]=ACK  vote[2]=ACK",
    "status = PROPOSED",
]

CL_COMMITTED = [
    "proposer = 0    epoch = 42",
    "old_slot = 0x0    new_slot = A",
    "vote[0]=ACK  vote[1]=ACK  vote[2]=ACK",
    "status = COMMITTED",
]


# ================================================================
#  STEP 1: Write KV to local DRAM + stage in CXL
# ================================================================
slide = _add_slide("Step 1: Write KV to local DRAM + stage in CXL Staging Buffer")
_draw_node(slide, 0, "Node 0 (Proposer)", COMP, COMP_L, kv_text="hello:world", hi_kv=True)
_draw_node(slide, 1, "Node 1 (Voter)", VOTER, VOTER_L)
_draw_node(slide, 2, "Node 2 (Voter)", VOTER, VOTER_L)
_draw_cxl(slide, CL_EMPTY, "op_status = ?", "hello:world", hi_st=True)
_switch_lines(slide)
# arrow: Node 0 -> local KV (internal, just label)
# arrow: Node 0 -> Staging
_add_arrow(slide, COL_X[0]+NW/2+0.5, NY+NH, ST_X+ST_W/2, CXL_Y, COMP, Pt(2))
_add_text(slide, 3.5, 3.2, 4, 0.5,
          "Node 0 writes KV to own DRAM\n+ copies to CXL Staging for replication", 9, COMP, bold=True)
_legend(slide)

# ================================================================
#  STEP 2: Write OpLog = IN_PROGRESS
# ================================================================
slide = _add_slide("Step 2: Write OpLog (status = IN_PROGRESS)")
_draw_node(slide, 0, "Node 0 (Proposer)", COMP, COMP_L, kv_text="hello:world")
_draw_node(slide, 1, "Node 1 (Voter)", VOTER, VOTER_L)
_draw_node(slide, 2, "Node 2 (Voter)", VOTER, VOTER_L)
_draw_cxl(slide, CL_EMPTY, "op_status =\nIN_PROGRESS", "hello:world", hi_ol=True)
_switch_lines(slide)
_add_arrow(slide, COL_X[0]+NW/2, NY+NH, OL_X+OL_W/2, CXL_Y, COMP, Pt(2))
_add_text(slide, 2.0, 3.2, 5, 0.5,
          "Node 0 writes OpLog entry for crash recovery\nop_status := IN_PROGRESS", 9, COMP, bold=True)
_legend(slide)

# ================================================================
#  STEP 3: Claim + Propose (write ConsensusLog)
# ================================================================
slide = _add_slide("Step 3: Claim + Propose in ConsensusLog")
_draw_node(slide, 0, "Node 0 (Proposer)", COMP, COMP_L, kv_text="hello:world")
_draw_node(slide, 1, "Node 1 (Voter)", VOTER, VOTER_L)
_draw_node(slide, 2, "Node 2 (Voter)", VOTER, VOTER_L)
_draw_cxl(slide, CL_PROPOSED, "op_status =\nIN_PROGRESS", "hello:world", hi_cl=True)
_switch_lines(slide)
_add_arrow(slide, COL_X[0]+NW/2-0.3, NY+NH, CL_X+CL_W/2, CXL_Y, COMP, Pt(2.5))
_add_text(slide, 0.3, 3.0, 5, 0.8,
          "Node 0 writes ConsensusLog:\n"
          "  proposer=0, epoch=42\n"
          "  old=0x0, new=A\n"
          "  status := PROPOSED, vote[0] := ACK\n"
          "  + sfence",
          8, COMP, bold=True, align=PP_ALIGN.LEFT)

_add_text(slide, CL_X+CL_W+0.1, CL_Y+0.1, 2.5, 0.5,
          "FREE ->\nPROPOSED", 10, OK, bold=True)
_legend(slide)

# ================================================================
#  STEP 4: Voters poll CXL, check local, write ACK
# ================================================================
slide = _add_slide("Step 4: Voters poll ConsensusLog, check local DRAM, vote ACK")
_draw_node(slide, 0, "Node 0 (Proposer)", COMP, COMP_L, kv_text="hello:world")
_draw_node(slide, 1, "Node 1 (Voter)", VOTER, VOTER_L)
_draw_node(slide, 2, "Node 2 (Voter)", VOTER, VOTER_L)
_draw_cxl(slide, CL_VOTED, "op_status =\nIN_PROGRESS", "hello:world", hi_cl=True)
_switch_lines(slide)

# Node 1: read CXL (dashed), write vote (solid)
_add_arrow(slide, CL_X+CL_W/2+0.5, CXL_Y, COL_X[1]+NW/2-0.3, NY+NH, VOTER, Pt(1.5), dashed=True)
_add_arrow(slide, COL_X[1]+NW/2+0.3, NY+NH, CL_X+CL_W/2+0.8, CXL_Y, VOTER, Pt(2))

# Node 2: read CXL (dashed), write vote (solid)
_add_arrow(slide, CL_X+CL_W/2+1.2, CXL_Y, COL_X[2]+NW/2-0.3, NY+NH, VOTER, Pt(1.5), dashed=True)
_add_arrow(slide, COL_X[2]+NW/2+0.3, NY+NH, CL_X+CL_W/2+1.5, CXL_Y, VOTER, Pt(2))

_add_text(slide, 5.0, 3.0, 5.5, 0.8,
          "Each voter fiber:\n"
          "  1. poll CXL: see status == PROPOSED\n"
          "  2. check own local Slot X == 0x0 (matches old_slot)\n"
          "  3. store vote[i] := ACK + sfence",
          8, VOTER, bold=True, align=PP_ALIGN.LEFT)

_add_text(slide, CL_X+CL_W+0.1, CL_Y+0.8, 2.5, 0.5,
          "vote = [ACK,\n ACK, ACK]", 10, OK, bold=True)
_legend(slide)

# ================================================================
#  STEP 5: Proposer counts votes -> COMMITTED (COMMIT POINT)
# ================================================================
slide = _add_slide("Step 5: Count votes >= 2/3 -> status := COMMITTED  ==  COMMIT POINT",
                   COMMIT)
_draw_node(slide, 0, "Node 0 (Proposer)", COMP, COMP_L, kv_text="hello:world")
_draw_node(slide, 1, "Node 1 (Voter)", VOTER, VOTER_L)
_draw_node(slide, 2, "Node 2 (Voter)", VOTER, VOTER_L)
_draw_cxl(slide, CL_COMMITTED, "op_status =\nIN_PROGRESS", "hello:world", hi_cl=True)
_switch_lines(slide)

# Node 0 reads votes (dashed), writes COMMITTED (solid)
_add_arrow(slide, CL_X+CL_W/2-0.5, CXL_Y, COL_X[0]+NW/2+0.3, NY+NH, COMMIT, Pt(1.5), dashed=True)
_add_arrow(slide, COL_X[0]+NW/2-0.3, NY+NH, CL_X+CL_W/2-0.8, CXL_Y, COMMIT, Pt(2.5))

_add_text(slide, 0.3, 3.0, 4.5, 0.6,
          "Node 0: poll votes -> 3 ACK >= 2\n"
          "store status := COMMITTED + sfence",
          9, COMMIT, bold=True, align=PP_ALIGN.LEFT)

_add_text(slide, CL_X+CL_W+0.1, CL_Y+0.1, 2.5, 0.5,
          "PROPOSED ->\nCOMMITTED", 10, COMMIT, bold=True)

# commit point line
line = slide.shapes.add_connector(1, _in(0.3), _in(3.85), _in(13.0), _in(3.85))
line.line.color.rgb = COMMIT
line.line.width = Pt(3)
_add_text(slide, 9.5, 3.6, 3.5, 0.35, "COMMIT POINT", 14, COMMIT, bold=True)
_legend(slide)

# ================================================================
#  STEP 6: All nodes apply to local DRAM
# ================================================================
slide = _add_slide("Step 6: All nodes see COMMITTED -> apply to local DRAM", OK)
_draw_node(slide, 0, "Node 0 (Proposer)", COMP, COMP_L,
           idx_text="Slot X = A", kv_text="hello:world", hi_idx=True)
_draw_node(slide, 1, "Node 1 (Voter)", VOTER, VOTER_L,
           idx_text="Slot X = A", kv_text="hello:world", hi_idx=True, hi_kv=True)
_draw_node(slide, 2, "Node 2 (Voter)", VOTER, VOTER_L,
           idx_text="Slot X = A", kv_text="hello:world", hi_idx=True, hi_kv=True)
_draw_cxl(slide, CL_COMMITTED, "op_status =\nIN_PROGRESS", "hello:world")
_switch_lines(slide)

# arrows from CXL up to all nodes
for col in range(3):
    _add_arrow(slide, CL_X+CL_W/2, CXL_Y, COL_X[col]+NW/2, NY+NH,
               OK, Pt(1.5), dashed=True)
# arrows from staging to Node 1, 2
for col in [1, 2]:
    _add_arrow(slide, ST_X+ST_W/2, CXL_Y, COL_X[col]+NW/2+0.5, NY+NH,
               OK, Pt(1.5), dashed=True)

_add_text(slide, 4.0, 3.0, 6, 0.8,
          "Each node's applier fiber:\n"
          "  1. see status == COMMITTED in ConsensusLog\n"
          "  2. write Slot X := A to own Local Index\n"
          "  3. copy KV from CXL Staging to own Local KV Data",
          8, OK, bold=True, align=PP_ALIGN.LEFT)
_legend(slide)

# ================================================================
#  STEP 7: Update OpLog = COMMITTED
# ================================================================
slide = _add_slide("Step 7: Update OpLog (status = COMMITTED)")
_draw_node(slide, 0, "Node 0 (Proposer)", COMP, COMP_L,
           idx_text="Slot X = A", kv_text="hello:world")
_draw_node(slide, 1, "Node 1 (Voter)", VOTER, VOTER_L,
           idx_text="Slot X = A", kv_text="hello:world")
_draw_node(slide, 2, "Node 2 (Voter)", VOTER, VOTER_L,
           idx_text="Slot X = A", kv_text="hello:world")
_draw_cxl(slide, CL_COMMITTED, "op_status =\nCOMMITTED", "hello:world", hi_ol=True)
_switch_lines(slide)

_add_arrow(slide, COL_X[0]+NW/2, NY+NH, OL_X+OL_W/2, CXL_Y, COMP, Pt(2))
_add_text(slide, 2.0, 3.2, 4.5, 0.4,
          "Node 0 updates OpLog: op_status := COMMITTED", 9, COMP, bold=True)

# final state box
_add_box(slide, 2.0, CXL_Y+CXL_H+0.15, 9.5, 0.4,
         RGBColor(0xE8, 0xF5, 0xE9), OK, Pt(1.5),
         "Final: All 3 nodes have Slot X = A | KV replicated | OpLog committed", 9, OK)
_legend(slide)

# ── Save ──
out = "docs/cxl_consensus_steps.pptx"
prs.save(out)
print(f"Saved {out}")
