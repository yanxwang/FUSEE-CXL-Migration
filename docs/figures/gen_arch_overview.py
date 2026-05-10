"""
Protocol A v2 — overall architecture (single PPT slide).

Visual conventions:
  - threads          : ellipses  (worker / forwarder / dispatcher)
  - data structures  : rectangles, sized by relative footprint:
        DRAM   :  KvCachePool > SlotDirectory > ShardingTable
        CXL    :  Hashtable (flat, on top, with visible bucket cells)
                  KV Blockpool (tall, on bottom, with H-segment partition)
  - SPSC rings act as bridges connecting the two DRAMs across CXL.

Layout:
  Top    :  Host 0 DRAM             |    Host 1 DRAM
            (data structures stacked top-down by size, threads on bottom row)
  Middle :  ForwardRing  /  InvalRing      (bridges, on CXL)
  Bottom :  Hashtable  (flat metadata band)
            KV Blockpool (tall data tier, with per-host segments)

Source of truth:
  docs/design_goals.md, docs/protocol_a_architecture_blueprint.md.

Run:
  python3 gen_arch_overview.py
Output:
  arch_overview.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree


# ----------- palette -----------
COL_DRAM_BG     = RGBColor(0xFD, 0xF2, 0xE9)
COL_DRAM_BORDER = RGBColor(0xC0, 0x55, 0x21)
COL_CXL_BG      = RGBColor(0xE8, 0xF1, 0xFA)
COL_CXL_BORDER  = RGBColor(0x1F, 0x4E, 0x79)

# DRAM data tiles — orange family, sized by importance
COL_KVC         = RGBColor(0xF6, 0xB2, 0x6B)   # KvCachePool (largest, deeper)
COL_KVC_BD      = RGBColor(0xB4, 0x5F, 0x06)
COL_SDIR        = RGBColor(0xFC, 0xD5, 0xA5)   # SlotDirectory
COL_SDIR_BD     = RGBColor(0xC0, 0x6B, 0x0A)
COL_STBL        = RGBColor(0xFD, 0xE9, 0xCB)   # ShardingTable (smallest, palest)
COL_STBL_BD     = RGBColor(0xC4, 0x82, 0x32)

# CXL data tiles — blue family
COL_HT          = RGBColor(0xCF, 0xE2, 0xF3)
COL_HT_BD       = RGBColor(0x35, 0x6E, 0xA8)
COL_HT_CELL     = RGBColor(0xE6, 0xEE, 0xF8)   # bucket inner cells
COL_POOL        = RGBColor(0xCF, 0xE2, 0xF3)
COL_POOL_BD     = RGBColor(0x35, 0x6E, 0xA8)
COL_POOL_SEG    = RGBColor(0xDD, 0xE9, 0xF6)   # segment inner box
COL_POOL_BLK    = RGBColor(0xA8, 0xC4, 0xE3)   # mini value blocks

# Ring channels
COL_FRING       = RGBColor(0xC2, 0xE0, 0xB4)
COL_FRING_BD    = RGBColor(0x38, 0x76, 0x1D)
COL_IRING       = RGBColor(0xF1, 0xC2, 0xC2)
COL_IRING_BD    = RGBColor(0xCC, 0x00, 0x00)

# Threads (ellipses) — distinct colors
COL_WKR         = RGBColor(0xCF, 0xDF, 0xF6)   # Workers — soft blue
COL_WKR_BD      = RGBColor(0x1F, 0x4E, 0x79)
COL_FRESP       = RGBColor(0xD9, 0xEA, 0xD3)   # Responder — soft green
COL_FRESP_BD    = RGBColor(0x38, 0x76, 0x1D)
COL_DISP        = RGBColor(0xF4, 0xCC, 0xCC)   # Dispatcher — soft red
COL_DISP_BD     = RGBColor(0xCC, 0x00, 0x00)

COL_TEXT        = RGBColor(0x1F, 0x1F, 0x1F)
COL_AR_FLUSH    = RGBColor(0x1F, 0x4E, 0x79)
COL_AR_FWD      = RGBColor(0x38, 0x76, 0x1D)
COL_AR_INV      = RGBColor(0xCC, 0x00, 0x00)


# ----------- helpers -----------

def _set_text(tf, lines, size=11, bold=False, italic=False,
              align=PP_ALIGN.CENTER, color=COL_TEXT,
              margin_em=18000):
    tf.word_wrap = True
    tf.margin_left = Emu(margin_em)
    tf.margin_right = Emu(margin_em)
    tf.margin_top = Emu(margin_em)
    tf.margin_bottom = Emu(margin_em)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color


def _rect(slide, x, y, w, h, fill, border, text=None,
          size=11, bold=True, line_w=1.0, rounded=True):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shp_type, x, y, w, h)
    if rounded:
        shp.adjustments[0] = 0.10
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border
    shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if text is not None:
        _set_text(shp.text_frame, text, size=size, bold=bold)
    else:
        shp.text_frame.text = ""
    return shp


def _ellipse(slide, x, y, w, h, fill, border, text,
             size=10, bold=True, line_w=1.2):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = border
    shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    _set_text(shp.text_frame, text, size=size, bold=bold)
    return shp


def _label(slide, x, y, w, h, text, size=10, bold=False, italic=False,
           align=PP_ALIGN.CENTER, color=COL_TEXT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    _set_text(tb.text_frame, text, size=size, bold=bold, italic=italic,
              align=align, color=color)
    return tb


def _arrow(slide, x1, y1, x2, y2, color, dbl=False, line_w=1.8,
           dash=False):
    shp = slide.shapes.add_connector(1, x1, y1, x2, y2)
    line = shp.line
    line.color.rgb = color
    line.width = Pt(line_w)
    ln = shp.line._get_or_add_ln()
    for tag in ('tailEnd', 'headEnd'):
        existing = ln.find(qn('a:' + tag))
        if existing is not None:
            ln.remove(existing)
    head = etree.SubElement(ln, qn('a:headEnd'))
    head.set('type', 'triangle')
    head.set('w', 'med')
    head.set('len', 'med')
    if dbl:
        tail = etree.SubElement(ln, qn('a:tailEnd'))
        tail.set('type', 'triangle')
        tail.set('w', 'med')
        tail.set('len', 'med')
    if dash:
        prstDash = etree.SubElement(ln, qn('a:prstDash'))
        prstDash.set('val', 'dash')
    return shp


# ---------- specialised renderers ----------

def _hashtable_band(slide, x, y, w, h, n_buckets=16):
    """Flat band with visible bucket cells in a row."""
    _rect(slide, x, y, w, h, COL_HT, COL_HT_BD, text=None,
          line_w=1.5, rounded=False)
    label_w = Inches(1.20)
    _label(slide, x + Inches(0.12), y + (h - Inches(0.25)) / 2,
           label_w, Inches(0.25),
           "Hashtable", size=12, bold=True, align=PP_ALIGN.LEFT,
           color=COL_HT_BD)
    cell_x0 = x + label_w + Inches(0.18)
    cell_x1 = x + w - Inches(0.18)
    cell_w = (cell_x1 - cell_x0 - (n_buckets - 1) * Inches(0.04)) / n_buckets
    cell_h = h - Inches(0.18)
    cell_y = y + Inches(0.09)
    cx = cell_x0
    for _ in range(n_buckets):
        cell = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, cx, cell_y, cell_w, cell_h)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COL_HT_CELL
        cell.line.color.rgb = COL_HT_BD
        cell.line.width = Pt(0.4)
        cell.shadow.inherit = False
        cell.text_frame.text = ""
        cx += cell_w + Inches(0.04)


def _blockpool_block(slide, x, y, w, h, n_segments=2):
    """Tall pool, partitioned into H owner-private segments separated by a
    visible barrier line (no cross-host writes)."""
    _rect(slide, x, y, w, h, COL_POOL, COL_POOL_BD, text=None,
          line_w=1.5, rounded=False)
    label_w = Inches(1.50)
    _label(slide, x + Inches(0.12), y + Inches(0.12),
           label_w, Inches(0.30),
           "KV Blockpool", size=12, bold=True, align=PP_ALIGN.LEFT,
           color=COL_POOL_BD)

    seg_x0 = x + label_w + Inches(0.18)
    seg_x1 = x + w - Inches(0.18)
    seg_y  = y + Inches(0.18)
    seg_h  = h - Inches(0.30)
    gap = Inches(0.18)
    seg_w = (seg_x1 - seg_x0 - (n_segments - 1) * gap) / n_segments

    # subtle per-host tint so the two halves look like distinct owner regions
    seg_fills = [
        RGBColor(0xDD, 0xE9, 0xF6),   # host 0 — bluish
        RGBColor(0xE9, 0xE0, 0xF1),   # host 1 — purplish
    ]

    sx = seg_x0
    seg_centers = []
    for i in range(n_segments):
        seg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, sx, seg_y, seg_w, seg_h)
        seg.fill.solid()
        seg.fill.fore_color.rgb = seg_fills[i % len(seg_fills)]
        seg.line.color.rgb = COL_POOL_BD
        seg.line.width = Pt(0.8)
        seg.shadow.inherit = False
        seg.text_frame.text = ""
        # segment label — "owner-private" wording makes ownership explicit
        _label(slide, sx + Inches(0.10), seg_y + Inches(0.04),
               seg_w - Inches(0.20), Inches(0.22),
               f"Host {i} segment   (owner-private)",
               size=9, italic=True, bold=True,
               align=PP_ALIGN.LEFT, color=COL_POOL_BD)
        # variable-size mini blocks — first row only (kept compact)
        block_y = seg_y + Inches(0.34)
        block_h = Inches(0.28)
        block_widths = [0.55, 1.05, 0.40, 0.85, 0.65, 0.35]
        bx = sx + Inches(0.12)
        max_x = sx + seg_w - Inches(0.12)
        row = 0
        for bw in block_widths:
            bw_in = Inches(bw)
            if bx + bw_in > max_x:
                row += 1
                bx = sx + Inches(0.12)
                if row >= 2:
                    break
            blk = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                bx, block_y + Inches(row * 0.34),
                bw_in, block_h)
            blk.fill.solid()
            blk.fill.fore_color.rgb = COL_POOL_BLK
            blk.line.color.rgb = COL_POOL_BD
            blk.line.width = Pt(0.4)
            blk.shadow.inherit = False
            blk.text_frame.text = ""
            bx += bw_in + Inches(0.06)
        seg_centers.append(sx + seg_w / 2)
        sx += seg_w + gap

    # ---- partition barrier between segments (visible thick dashed line)
    if n_segments == 2:
        bar_x = (seg_centers[0] + seg_centers[1]) / 2
        bar_top = seg_y + Inches(0.04)
        bar_bot = seg_y + seg_h - Inches(0.04)
        bar = slide.shapes.add_connector(1, bar_x, bar_top,
                                         bar_x, bar_bot)
        bar.line.color.rgb = COL_POOL_BD
        bar.line.width = Pt(2.5)
        ln = bar.line._get_or_add_ln()
        prst = etree.SubElement(ln, qn('a:prstDash'))
        prst.set('val', 'dash')
        # tiny "no cross-host writes" tag at the bar's middle
        tag_w = Inches(1.20)
        tag_x = bar_x - tag_w / 2
        tag_y = (bar_top + bar_bot) / 2 - Inches(0.10)
        _label(slide, tag_x, tag_y, tag_w, Inches(0.20),
               "no cross-host writes",
               size=7, italic=True, bold=True,
               align=PP_ALIGN.CENTER, color=COL_POOL_BD)


# ---------- main ----------

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    _label(slide, Inches(0.30), Inches(0.10), Inches(12.7), Inches(0.42),
           "Protocol A v2 — Overall Architecture",
           size=20, bold=True, align=PP_ALIGN.LEFT)

    # ========================================================
    # TOP — Two host DRAMs
    # ========================================================
    dram_y, dram_h = Inches(0.62), Inches(2.95)
    dram_w = Inches(6.00)
    h0_x = Inches(0.30)
    h1_x = Inches(7.03)        # 0.73 inch gap

    anchors = {}

    for hx, hid in [(h0_x, 0), (h1_x, 1)]:
        # outer DRAM box
        outer = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, hx, dram_y, dram_w, dram_h)
        outer.adjustments[0] = 0.04
        outer.fill.solid()
        outer.fill.fore_color.rgb = COL_DRAM_BG
        outer.line.color.rgb = COL_DRAM_BORDER
        outer.line.width = Pt(2.0)
        outer.shadow.inherit = False
        outer.text_frame.text = ""

        _label(slide, hx + Inches(0.15), dram_y + Inches(0.05),
               Inches(5.0), Inches(0.25),
               f"Host {hid} — DRAM (MAP_SHARED)",
               size=11, bold=True, align=PP_ALIGN.LEFT,
               color=COL_DRAM_BORDER)

        # ---- DRAM data structures (sized by importance, top-down) ----
        # KvCachePool — biggest
        kvc_w = dram_w - Inches(0.30)
        kvc_h = Inches(0.75)
        kvc_x = hx + (dram_w - kvc_w) / 2
        kvc_y = dram_y + Inches(0.38)
        _rect(slide, kvc_x, kvc_y, kvc_w, kvc_h,
              COL_KVC, COL_KVC_BD, "KvCachePool",
              size=13, bold=True, line_w=1.2)

        # SlotDirectory — medium
        sd_w = Inches(4.20)
        sd_h = Inches(0.50)
        sd_x = hx + (dram_w - sd_w) / 2
        sd_y = kvc_y + kvc_h + Inches(0.10)
        _rect(slide, sd_x, sd_y, sd_w, sd_h,
              COL_SDIR, COL_SDIR_BD, "SlotDirectory",
              size=11, bold=True, line_w=1.0)

        # ShardingTable — smallest
        st_w = Inches(2.50)
        st_h = Inches(0.32)
        st_x = hx + (dram_w - st_w) / 2
        st_y = sd_y + sd_h + Inches(0.10)
        _rect(slide, st_x, st_y, st_w, st_h,
              COL_STBL, COL_STBL_BD, "ShardingTable",
              size=10, bold=True, line_w=0.8)

        # ---- threads (ellipses) along bottom of DRAM ----
        th_y = dram_y + dram_h - Inches(0.78)
        th_h = Inches(0.62)
        th_w = Inches(1.55)
        th_pad = (dram_w - th_w * 3) / 4
        tx = hx + th_pad
        wkr = _ellipse(slide, tx, th_y, th_w, th_h,
                       COL_WKR, COL_WKR_BD, "Workers × T",
                       size=10, bold=True)
        anchors[(hid, 'workers')] = (tx + th_w / 2, th_y + th_h)
        tx += th_w + th_pad
        fr = _ellipse(slide, tx, th_y, th_w, th_h,
                      COL_FRESP, COL_FRESP_BD,
                      ["Forward", "Responder"],
                      size=10, bold=True)
        anchors[(hid, 'fwdresp')] = (tx + th_w / 2, th_y + th_h)
        tx += th_w + th_pad
        cd = _ellipse(slide, tx, th_y, th_w, th_h,
                      COL_DISP, COL_DISP_BD,
                      ["Cache", "Dispatcher"],
                      size=10, bold=True)
        anchors[(hid, 'disp')] = (tx + th_w / 2, th_y + th_h)

    # ========================================================
    # BOTTOM — CXL region (rings + flat hashtable + tall blockpool)
    # ========================================================
    cxl_y = Inches(3.72)
    cxl_h = Inches(3.55)
    cxl_x = Inches(0.30)
    cxl_w = Inches(12.73)

    cxl_outer = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, cxl_x, cxl_y, cxl_w, cxl_h)
    cxl_outer.adjustments[0] = 0.025
    cxl_outer.fill.solid()
    cxl_outer.fill.fore_color.rgb = COL_CXL_BG
    cxl_outer.line.color.rgb = COL_CXL_BORDER
    cxl_outer.line.width = Pt(2.0)
    cxl_outer.shadow.inherit = False
    cxl_outer.text_frame.text = ""

    _label(slide, cxl_x + Inches(0.15), cxl_y + Inches(0.04),
           Inches(8.0), Inches(0.26),
           "CXL shared region   (single physical copy across hosts)",
           size=11, bold=True, align=PP_ALIGN.LEFT, color=COL_CXL_BORDER)

    # --- bridge rings (top of CXL, span between DRAMs) ---
    bridge_l = Inches(1.25)
    bridge_r = Inches(12.05)
    bridge_w = bridge_r - bridge_l

    fr_y = cxl_y + Inches(0.40)
    fr_h = Inches(0.45)
    _rect(slide, bridge_l, fr_y, bridge_w, fr_h,
          COL_FRING, COL_FRING_BD, "ForwardRing  [H][H]   (SPSC)",
          size=11, bold=True, line_w=1.2)
    anchors[('fwdring', 'top_l')] = (bridge_l + Inches(0.50), fr_y)
    anchors[('fwdring', 'top_r')] = (bridge_r - Inches(0.50), fr_y)

    ir_y = fr_y + fr_h + Inches(0.14)
    ir_h = Inches(0.45)
    _rect(slide, bridge_l, ir_y, bridge_w, ir_h,
          COL_IRING, COL_IRING_BD, "InvalRing  [H][H]   (SPSC)",
          size=11, bold=True, line_w=1.2)
    anchors[('invring', 'top_l')] = (bridge_l + Inches(0.50), ir_y)
    anchors[('invring', 'top_r')] = (bridge_r - Inches(0.50), ir_y)

    # --- flat Hashtable band (metadata, on top of pool) ---
    ht_y = ir_y + ir_h + Inches(0.30)
    ht_h = Inches(0.45)
    _hashtable_band(slide,
                    cxl_x + Inches(0.18), ht_y,
                    cxl_w - Inches(0.36), ht_h,
                    n_buckets=18)

    # --- tall KV Blockpool (data tier, below hashtable) ---
    pool_y = ht_y + ht_h + Inches(0.10)
    pool_h = cxl_y + cxl_h - pool_y - Inches(0.18)
    _blockpool_block(slide,
                     cxl_x + Inches(0.18), pool_y,
                     cxl_w - Inches(0.36), pool_h,
                     n_segments=2)

    # ========================================================
    # ARROWS
    # ========================================================

    # forward path: H0 Workers -> FwdRing left -> H1 ForwardResponder
    x0w, y0w = anchors[(0, 'workers')]
    x_fl, y_fl = anchors[('fwdring', 'top_l')]
    _arrow(slide, x0w, y0w, x_fl, y_fl, COL_AR_FWD, line_w=2.0)

    x_fr, y_fr = anchors[('fwdring', 'top_r')]
    x1fr, y1fr = anchors[(1, 'fwdresp')]
    _arrow(slide, x_fr, y_fr, x1fr, y1fr, COL_AR_FWD, line_w=2.0)

    # invalidate path: H0 Workers -> InvalRing -> H1 CacheDispatcher
    x_il, y_il = anchors[('invring', 'top_l')]
    _arrow(slide,
           x0w + Inches(0.20), y0w,
           x_il, y_il,
           COL_AR_INV, line_w=2.0)
    x_ir, y_ir = anchors[('invring', 'top_r')]
    x1d, y1d = anchors[(1, 'disp')]
    _arrow(slide,
           x_ir, y_ir,
           x1d - Inches(0.20), y1d,
           COL_AR_INV, line_w=2.0)

    # direct CXL data access (flush+fence) - one per host, on outer edges
    a_x = h0_x + Inches(0.40)
    _arrow(slide, a_x, dram_y + dram_h,
           a_x, ht_y + ht_h / 2,
           COL_AR_FLUSH, dbl=True, line_w=1.6)
    _label(slide, a_x + Inches(0.05),
           dram_y + dram_h + Inches(0.10),
           Inches(2.6), Inches(0.20),
           "load/store + clflushopt + sfence",
           size=8, italic=True, align=PP_ALIGN.LEFT, color=COL_AR_FLUSH)

    b_x = h1_x + dram_w - Inches(0.40)
    _arrow(slide, b_x, dram_y + dram_h,
           b_x, ht_y + ht_h / 2,
           COL_AR_FLUSH, dbl=True, line_w=1.6)
    _label(slide, b_x - Inches(2.65),
           dram_y + dram_h + Inches(0.10),
           Inches(2.6), Inches(0.20),
           "load/store + clflushopt + sfence",
           size=8, italic=True, align=PP_ALIGN.RIGHT, color=COL_AR_FLUSH)

    # ========================================================
    # legend (top-right of title bar)
    # ========================================================

    def _swatch(x, y, color, label, dx_label=Inches(0.20)):
        sw = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y,
                                    Inches(0.16), Inches(0.10))
        sw.fill.solid()
        sw.fill.fore_color.rgb = color
        sw.line.color.rgb = color
        sw.line.width = Pt(0.5)
        sw.shadow.inherit = False
        _label(slide, x + dx_label, y - Inches(0.05),
               Inches(2.6), Inches(0.20),
               label, size=8.5, align=PP_ALIGN.LEFT,
               color=RGBColor(0x44, 0x44, 0x44))

    leg_y = Inches(0.20)
    _swatch(Inches(7.10), leg_y, COL_AR_FWD,  "write-forward")
    _swatch(Inches(8.90), leg_y, COL_AR_INV,  "invalidate")
    _swatch(Inches(10.45), leg_y, COL_AR_FLUSH,
            "direct CXL (clflushopt+sfence)")

    out = "/home/yanwang/FUSEE/docs/figures/arch_overview.pptx"
    prs.save(out)
    print(f"saved: {out}")


if __name__ == "__main__":
    main()
